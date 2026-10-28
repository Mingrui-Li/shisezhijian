"""Build copyright filing drafts from current client source and reviewed Markdown.

Run with the Codex bundled Python environment (python-docx required).
Rendering and visual verification are separate required steps.
"""
from pathlib import Path
import csv
import hashlib
import json
import re
import shutil
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output/copyright'
OUT.mkdir(parents=True, exist_ok=True)
NAME = '拾色之间游戏软件'
VERSION = 'V1.1.0'
OWNER = '济南青骓互动信息科技有限公司'
BASE = f'{NAME}_{VERSION}'


def font(run, size=11, family='Arial', east='Arial Unicode MS', bold=False):
    run.font.name = family
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), east)
    for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
        run._element.get_or_add_rPr().rFonts.attrib.pop(qn('w:' + attr), None)


def field(p, name):
    r = p.add_run()
    font(r, 9)
    f = OxmlElement('w:fldSimple')
    f.set(qn('w:instr'), name)
    r._r.addnext(f)


def document(kind, code=False):
    doc = Document()
    s = doc.sections[0]
    s.page_width, s.page_height = Mm(210), Mm(297)
    s.left_margin = s.right_margin = Mm(15 if code else 20)
    s.top_margin, s.bottom_margin = Mm(16 if code else 19), Mm(17)
    s.header_distance, s.footer_distance = Mm(7), Mm(8)
    for style_name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        sty = doc.styles[style_name]
        sty.font.name = 'Arial'
        sty.font.color.rgb = RGBColor(0, 0, 0)
        sty.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial Unicode MS')
        for attr in ['asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme']:
            sty.element.get_or_add_rPr().rFonts.attrib.pop(qn('w:' + attr), None)
        for border in sty.element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
    normal = doc.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(18)
    normal.paragraph_format.widow_control = False
    for name, size in [('Title', 20), ('Heading 1', 16), ('Heading 2', 12)]:
        sty = doc.styles[name]
        sty.font.size = Pt(size)
        sty.font.bold = True
        sty.paragraph_format.space_before = Pt(0 if name != 'Heading 2' else 5)
        sty.paragraph_format.space_after = Pt(7 if name != 'Heading 2' else 4)
        sty.paragraph_format.line_spacing = 1.1
    h = s.header.paragraphs[0]
    doc.styles['Header'].paragraph_format.tab_stops.clear_all()
    h.paragraph_format.tab_stops.clear_all()
    h.paragraph_format.tab_stops.add_tab_stop(Mm(180 if code else 170), WD_ALIGN_PARAGRAPH.RIGHT)
    font(h.add_run(f'{NAME} {VERSION}'), 9)
    font(h.add_run('\t第 '), 9)
    field(h, 'PAGE')
    font(h.add_run(' 页'), 9)
    f = s.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(f.add_run(OWNER), 8.5)
    doc.core_properties.title = f'{NAME} {VERSION} {kind}'
    doc.core_properties.author = OWNER
    doc.core_properties.subject = '计算机软件著作权登记申请准备材料'
    return doc


def body(doc, text, size=11):
    p = doc.add_paragraph()
    font(p.add_run(text), size)
    return p


def table(doc, rows):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    t.columns[0].width, t.columns[1].width = Mm(39), Mm(131)
    for idx, vals in enumerate(rows):
        cells = t.add_row().cells
        cells[0].width, cells[1].width = Mm(39), Mm(131)
        for j, val in enumerate(vals):
            c = cells[j]
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = c._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right']:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9')
                borders.append(e)
            pr.append(borders)
            margins = OxmlElement('w:tcMar')
            for side, value in [('top', '70'), ('bottom', '70'), ('left', '100'), ('right', '100')]:
                e = OxmlElement('w:' + side); e.set(qn('w:w'), value); e.set(qn('w:type'), 'dxa'); margins.append(e)
            pr.append(margins)
            shading = OxmlElement('w:shd')
            shading.set(qn('w:fill'), '404040' if idx == 0 else ('F5F5F5' if idx % 2 == 0 else 'FFFFFF'))
            pr.append(shading)
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = Pt(15)
            r = p.add_run(val); font(r, 10.5, bold=idx == 0)
            if idx == 0: r.font.color.rgb = RGBColor(255, 255, 255)
        trpr = cells[0]._tc.getparent().get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trpr.append(cant)
        if idx == 0:
            repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
    return t


def from_markdown(source, suffix, manual=False):
    text = source.read_text()
    doc = document(suffix)
    for pi, page in enumerate(text.split('<!-- page -->')):
        if pi: doc.add_page_break()
        lines = page.strip().splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].strip()
            if not line: i += 1; continue
            if line.startswith('|'):
                rows = []
                while i < len(lines) and lines[i].strip().startswith('|'):
                    vals = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                    if not all(re.fullmatch('[-: ]+', v) for v in vals): rows.append(vals)
                    i += 1
                table(doc, rows)
                continue
            if line.startswith('# '):
                doc.add_paragraph(line[2:], 'Title' if pi == 0 else 'Heading 1')
            elif line.startswith('## '):
                doc.add_paragraph(line[3:], 'Heading 2')
            elif line.startswith('!['):
                image_path = re.search(r'\]\((.*?)\)', line).group(1)
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = 1
                p.paragraph_format.space_after = Pt(7)
                p.add_run().add_picture(str(ROOT / image_path), height=Mm(155))
            else:
                p = body(doc, line)
                if re.match(r'^图\d+', line):
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after = Pt(7)
                if re.match(r'^\d+\.\d+ ', line):
                    for r in p.runs: r.font.bold = True
            i += 1
    path = OUT / f'{BASE}_{suffix}.docx'
    doc.save(path)
    return path


FILES = [
    'Boot.ts', 'core/index.ts', 'application/Progress.ts',
    'application/RestoreProgress.ts', 'application/RuntimeConfig.ts',
    'platform/LocalSave.ts', 'platform/SnapshotStore.ts',
    'platform/RewardedVideo.ts', 'platform/Feedback.ts',
    'platform/BrowserTestBridge.ts', 'render/BoardLayout.ts',
    'render/LoadingScreen.ts', 'render/TutorialScreen.ts',
    'render/WinPanel.ts', 'theme/Theme.ts',
]


def source_document():
    srcroot = ROOT / 'game/assets/scripts'
    assert set(FILES) == {str(p.relative_to(srcroot)) for p in srcroot.rglob('*.ts')}
    snapshot = OUT / '核验附件/客户端源文件快照'
    snapshot.mkdir(parents=True, exist_ok=True)
    manifest, rows, original_total = [], [], 0
    full_text = []
    for rel in FILES:
        src = srcroot / rel
        data = src.read_bytes()
        original = data.decode('utf-8').splitlines()
        original_total += len(original)
        dest = snapshot / rel; dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(src, dest)
        manifest.append({'file': f'game/assets/scripts/{rel}', 'physicalLines': len(original),
            'nonblankLines': sum(bool(s.strip()) for s in original), 'sha256': hashlib.sha256(data).hexdigest()})
        full_text.append(f'// FILE: game/assets/scripts/{rel}\n' + data.decode('utf-8'))
        for original_line, code in enumerate(original, 1):
            if code.strip(): rows.append((rel, original_line, code))
    doc = document('源程序鉴别材料', code=True)
    page_count = (len(rows) + 49) // 50
    mapping = []
    for page in range(page_count):
        if page: doc.add_page_break()
        for row_no, (rel, original_line, code) in enumerate(rows[page*50:(page+1)*50], 1):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = Pt(9.7)
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_together = True
            p.paragraph_format.widow_control = False
            font(p.add_run(code.expandtabs(4)), 8.5, family='Liberation Mono')
            mapping.append([page+1, row_no, f'game/assets/scripts/{rel}', original_line])
    path = OUT / f'{BASE}_源程序鉴别材料.docx'
    doc.save(path)
    annex = OUT / '核验附件'
    with (annex / '源程序页行映射.csv').open('w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f); writer.writerow(['提交页码', '本页源码序号', '源文件', '原文件行号']); writer.writerows(mapping)
    (annex / '客户端源程序全文.txt').write_text('\n'.join(full_text), encoding='utf-8')
    report = {'softwareName': NAME, 'version': VERSION, 'owner': OWNER,
        'completionDate': '2026-10-28', 'publication': '未发表',
        'files': manifest, 'physicalLines': original_total, 'nonblankLines': len(rows),
        'submissionPages': page_count, 'sourceLinesPerFullPage': 50,
        'lastPageSourceLines': len(rows) % 50 or 50,
        'policy': '交存客户端全部TypeScript源程序，按所列模块顺序连接。仅省略空白行；非空白源码完整保留。长行自动折行，不改变源文件快照。未包含引擎或第三方实现、测试、工具及生成后的构建代码。'}
    (annex / '源程序清单与校验.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return path, report


if __name__ == '__main__':
    src, report = source_document()
    outputs = [from_markdown(ROOT / 'docs/copyright/申请信息草稿.md', '申请信息草稿'),
        from_markdown(ROOT / 'docs/copyright/软件说明书.md', '软件说明书', True), src]
    print(json.dumps({'outputs': [str(p) for p in outputs],
        'source': {k: report[k] for k in ['physicalLines', 'nonblankLines', 'submissionPages', 'lastPageSourceLines']}}, ensure_ascii=False))
